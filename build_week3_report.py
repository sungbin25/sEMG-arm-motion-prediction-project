"""Build the Korean PDF and editable DOCX from saved experiment outputs."""
from pathlib import Path
import json
import csv
from xml.sax.saxutils import escape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image, Table, TableStyle
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pypdfium2

ROOT=Path(__file__).parent
OUT=ROOT/'reports/week3';OUT.mkdir(parents=True,exist_ok=True)
DATA=ROOT/'outputs/week3'
result=json.loads((DATA/'results.json').read_text(encoding='utf-8'))
config=json.loads((DATA/'config.json').read_text(encoding='utf-8'))
blocks=[]
def text(t,kind='body'):blocks.append((kind,t))
def table(rows,widths):blocks.append(('table',(rows,widths)))
def picture(path,w,h):blocks.append(('image',(str(path),w,h)))
def page():blocks.append(('page',None))
def pct(x):return f'{x*100:.2f}'
def pvalue(x):return '검정 불가' if x is None else f'{x:.6g}'
holdout=result['holdout'];cv=result['cv'];tests=result['paired_tests']
best=max(holdout,key=lambda r:r['f1_macro'])
worst=min(holdout,key=lambda r:r['f1_macro'])
text('Step2 3주차 결과보고서','title')
text('sEMG 모델 비교와 통계 검정','h1')
text('오성빈 | 2026-10-20','small')
text(f"SVM, Random Forest, XGBoost를 동일한 2주차 특징과 시행 분할로 평가했다. 홀드아웃 Macro F1이 가장 높은 모델은 {best['model']}({pct(best['f1_macro'])}%)였다. 5-fold 교차검증, 대응표본 t-검정, 혼동행렬과 실행 시간을 함께 확인했으며, 이 결과는 새로운 피험자에 대한 성능을 의미하지 않는다.")
text('1 실험 개요','h2')
text('8명 × 8조건의 64개 기록을 8클래스로 분류했다. 2주차의 시행 ID 분할을 유지하여 학습 1,536시행·29,184창과 테스트 384시행·7,296창을 사용했다. 각 창은 200 ms이며 50% 중첩으로 25개 특징을 추출했다. 공개 신호에 추가 필터를 적용하지 않았고, 분석 구간은 마커 기준 1.5~3.5초라는 잠정 설정을 유지했다.')
table([['모델','고정한 주요 설정'],['SVM','RBF, C=1.0, gamma=scale, decision_function_shape=ovr'],['RF','n_estimators=100, random_state=42, n_jobs=4'],['XGB','n_estimators=400, max_depth=6, learning_rate=0.05, subsample=0.9, colsample_bytree=0.8, random_state=42']],[65,442])
text('세 모델 모두 StandardScaler와 분류기를 연결한 Pipeline을 사용했다. XGBoost는 CPU hist 방식과 4개 작업 스레드를 명시했다. SVM 캐시는 1,024 MB이며 RF·XGB의 작업 스레드는 4개다. 하이퍼파라미터·분석 구간 탐색은 수행하지 않았다.','small')
text('2 홀드아웃 성능 비교','h2')
table([['모델','Accuracy','Precision','Recall','F1']]+[[r['model']]+[pct(r[k]) for k in ('accuracy','precision_macro','recall_macro','f1_macro')] for r in holdout],[65,110,110,110,112])
text('단위는 %이다. Precision·Recall·F1은 클래스별 지표의 macro 평균이며 강의 실습의 선택이다. 모든 클래스의 테스트 창 수는 912개로 같다. 수치는 윈도우 단위 평가이며 시행별 다수결 정확도가 아니다. 테스트 성능을 이용한 설정 재조정은 수행하지 않았다.','small')
page()
text('3 시행 단위 5 fold 교차검증','h1')
text('기존 학습 데이터의 고유 시행 ID 1,536개에 StratifiedKFold(5, shuffle=True, random_state=42)를 적용하고, 시행 ID를 해당 윈도우 인덱스로 확장했다. 모든 모델은 동일한 fold를 공유하며, 각 fold에서 StandardScaler를 학습 부분에만 fit했다. CSV 한 파일이 시행 30개를 포함하므로 파일을 단일 시행으로 간주하지 않았다.')
folds=list(csv.DictReader((DATA/'cv_scores.csv').open()))
table([['모델','Fold 1','Fold 2','Fold 3','Fold 4','Fold 5','평균 ± SD']]+[[r['model']]+[pct(float(next(v['accuracy'] for v in folds if v['model']==r['model'] and int(v['fold'])==i))) for i in range(1,6)]+[f"{pct(r['accuracy_mean'])} ± {pct(r['accuracy_std'])}"] for r in cv],[48,65,65,65,65,65,134])
text('정확도 단위는 %이며 SD는 강의 코드와 동일하게 ddof=0으로 계산한 5개 fold의 표준편차다. 평균의 신뢰구간이 아니다. 교차검증의 검증 시행은 한 번씩만 등장하고, 동일 시행의 창이 해당 fold의 학습·검증에 함께 들어가지 않음을 검증했다.','small')
picture(DATA/'cv_accuracy.png',470,199)
text('4 대응표본 t 검정','h2')
table([['비교','평균 차이 pp','t','원 p값','Holm 보정 p값']]+[[r['model_a']+' vs '+r['model_b'],f"{r['mean_accuracy_difference']*100:+.3f}",'-' if r['t'] is None else f"{r['t']:.3f}",pvalue(r['p_raw']),pvalue(r['p_holm'])] for r in tests],[90,95,65,125,132])
for r in tests:
    conclusion='검정 불가' if r['p_raw'] is None else '명목 유의수준 0.05에서 차이가 검출되었다' if r['p_raw']<.05 else '명목 유의수준 0.05에서 차이를 검출하지 못했다'
    text(f"{r['model_a']}와 {r['model_b']}: {conclusion}. 평균 차이는 앞 모델에서 뒤 모델을 뺀 값이다.",'small')
text('동일한 5개 fold의 정확도를 짝지었고 세 쌍의 비교에 Holm 보정도 제시했다. 다만 fold 간 학습 집합이 겹치고 관측 쌍은 5개뿐이므로 독립성·정규성 조건을 충분히 보장할 수 없다. 원 p값과 보정 p값 모두 탐색적 참고값이며, 보정만으로 fold 의존성이 해결되지는 않는다. p≥0.05는 동등성의 증거가 아니다.','small')
for name in ('SVM','RF','XGB'):
    page();e=result['errors'][name]
    text('5 혼동행렬과 오류 분석 '+name,'h1')
    picture(DATA/f'confusion_{name}.png',435,346)
    text('행은 실제 클래스, 열은 예측 클래스다. 모든 모델에 동일한 색 범위(0~912)를 적용했다. fist=주먹 쥐기, open=손 펴기, pronation=회내, supination=회외이며 30/60은 팔꿈치 각도다.','small')
    table([['클래스','Recall %']]+[[label,pct(value)] for label,value in zip(result['class_names'],e['class_recall'])],[245,262])
    text('가장 잘 분류된 클래스: '+', '.join(e['best_classes'])+'. 가장 낮은 재현율의 클래스: '+', '.join(e['worst_classes'])+'.','small')
    text('주요 오분류 방향: '+'; '.join(f"{r['actual']} → {r['predicted']} {r['count']}건" for r in e['top_errors'][:3])+'.','small')
    counts=e['error_types'];total=e['errors']
    text(f"총 오분류 {total}건 중 같은 동작·다른 각도 {counts['same_gesture_wrong_angle']}건, 같은 각도·다른 동작 {counts['same_angle_wrong_gesture']}건, 동작·각도 모두 다름 {counts['both_wrong']}건이다. 이는 실제 라벨 의미로 분류했으며 행렬에서 인접한 칸이라는 이유로 오류 종류를 정하지 않았다.",'small')
page()
text('6 계산 비용과 모델 선택','h1')
table([['모델','학습 시간 s','단일 창 평균 ms','단일 창 중앙값 ms','단일 창 P95 ms']]+[[r['model'],f"{r['fit_seconds']:.3f}",f"{r['single_sample_mean_ms']:.3f}",f"{r['single_sample_median_ms']:.3f}",f"{r['single_sample_p95_ms']:.3f}"] for r in holdout],[55,105,115,117,115])
table([['모델','7,296창 배치 추론 s','배치 추론 ms per window']]+[[r['model'],f"{r['batch_seconds']:.3f}",f"{r['batch_ms_per_window']:.4f}"] for r in holdout],[75,215,217])
text('학습 시간은 전체 학습 데이터에 대한 최종 Pipeline.fit 1회의 경과 시간이다. 단일 창 추론은 동일한 테스트 첫 창으로 10회 예열한 뒤 100회 predict 호출을 측정했다. 시간은 perf_counter로 기록했으며 표준화와 호출 비용을 포함한다. 배치 추론은 테스트 전체 1회에 대한 시간이다. 파일 읽기·특징 추출·센서 입력 시간은 포함하지 않았다.')
text('측정 환경','h2')
text(f"{config['platform']}; Python {config['python']}; 논리 CPU {config['logical_cpus']}개. CPU 식별 문자열: {config['processor']}. RF·XGB는 n_jobs=4이고 SVM은 같은 방식의 병렬 설정을 제공하지 않는다. GPU는 사용하지 않았다.")
text('패키지: '+', '.join(k+' '+v for k,v in config['versions'].items())+'.')
text('선택 기준','h2')
fastest=min(holdout,key=lambda r:r['single_sample_mean_ms'])
text(f"이번 홀드아웃 Macro F1 순위의 1위는 {best['model']}, 최하위는 {worst['model']}이다. 단일 창 평균 추론이 가장 빠른 모델은 {fastest['model']}이었다. 정확도와 지연시간의 우선순위에 따라 선택이 달라질 수 있으며, 이 순위가 다른 장비나 새로운 피험자에서도 유지된다고 단정할 수 없다.")
text('계산 비용의 해석','h2')
text('단일 창 호출은 실행 프레임워크와 스레드 관리 비용의 영향을 받을 수 있다. 배치 처리의 창당 시간과 실시간 단일 창 지연시간은 서로 다른 지표다. 학습 시간은 반복 실험의 평균이 아니며 OS 부하·캐시·스레드 수에 민감하다. 실제 배포 성능은 별도 환경에서 측정해야 한다.')
page()
text('7 최종 결과와 한계','h1')
text('실험에서 확인한 내용','h2')
text(f"같은 시행 분할과 특징을 사용한 세 모델의 홀드아웃 성능, 시행 단위 5-fold 변동, 오류 유형과 계산 비용을 산출했다. {best['model']}의 홀드아웃 Macro F1은 {pct(best['f1_macro'])}%이며, 5-fold 평균 정확도는 "+', '.join(f"{r['model']} {pct(r['accuracy_mean'])} ± {pct(r['accuracy_std'])}%" for r in cv)+'였다.')
text('강의에서 소개한 논문 결과와의 비교','h2')
rf=next(r for r in cv if r['model']=='RF')
text(f"강의자료는 논문의 RF 정확도를 96.72 ± 0.58%로 소개하고 세 모델 간 차이가 유의하지 않았다고 설명한다. 본 실험의 RF 교차검증 정확도는 {pct(rf['accuracy_mean'])} ± {pct(rf['accuracy_std'])}%였다. 강의에 제시되지 않은 논문의 다른 모델 수치나 세부 평가 조건을 추정해 채우지 않았다. 이 비교는 논문 수치의 정확한 재현을 확인하는 검증이 아니다.")
text('분석 시작점·마커 의미, 특징 정의, 사용 라이브러리·알고리즘 설정, 교차검증 모집단과 분할 절차의 차이가 결과 차이에 영향을 줄 수 있다. 특히 본 실험은 기존 홀드아웃을 보존하고 학습 부분만 교차검증했다. 어느 요인이 차이를 일으켰는지는 별도 통제 실험 없이 확정할 수 없다.')
text('해석의 한계','h2')
for value in [
    '1. 시행 단위 분할은 동일 피험자·동일 연속 기록의 다른 시행을 양쪽에 포함한다. 새 피험자 일반화는 LOSO 등 피험자 단위 평가로 별도 검증해야 한다.',
    '2. 1.5~3.5초는 잠정 분석 구간이다. 실제 동작 정렬 및 논문과 동일한 구간인지는 확인되지 않았고, 이번 실험에서도 구간 최적화를 수행하지 않았다.',
    '3. 중첩 창은 독립 표본이 아니다. 혼동행렬과 성능 지표는 윈도우 단위이며 창 수를 독립 시행 수로 해석할 수 없다.',
    '4. 5-fold 점수는 학습 집합을 공유하며 대응표본은 5개뿐이다. t-검정과 Holm 보정은 탐색적 결과이며 확정적인 우월성·동등성의 근거가 아니다.',
    '5. 하이퍼파라미터는 강의 설정으로 고정했다. 같은 데이터라도 모델별 탐색 예산을 맞춘 최적 성능 비교는 수행하지 않았다.',
    '6. 오분류의 생리적 원인은 혼동행렬만으로 판별할 수 없다. 유사한 근활성, 구간 정렬 또는 피험자 차이는 가능한 설명이며 추가 검증이 필요하다.'
]:text(value,'small')
text('재현과 근거 자료','h2')
text('실행: python -m pip install -r requirements-week3.txt → python -m unittest -v test_week2 test_week3 → python week3.py. 2주차 dataset.npz가 없으면 먼저 README의 원자료 다운로드 및 python week2.py 절차를 수행한다. outputs/week3에 fold별 결과, 시행 배정, 예측, 혼동행렬, 원시 시간 측정과 설정·입력 SHA-256을 저장했다.','small')
text('근거: 사용자 제공 Step2_3주차_강의와실습.pptx(특히 9~17쪽); 기존 2주차 결과보고서와 데이터셋; week3.py의 실제 실행 결과. 원자료 저장소: https://github.com/TestUser-alt-dev/sEMG-Hand-Gesture-Elbow-Angle-Dataset . 프로젝트: https://github.com/sungbin25/sEMG-arm-motion-prediction-project .','small')

# Native editable Word paragraphs and tables; figures remain images.
doc=Document();sec=doc.sections[0]
sec.page_width=Pt(595.28);sec.page_height=Pt(841.89)
sec.left_margin=sec.right_margin=Pt(44);sec.top_margin=Pt(40);sec.bottom_margin=Pt(45)
for name,size in [('Normal',9.5),('Title',24),('Heading 1',16),('Heading 2',11)]:
    st=doc.styles[name];st.font.name='Malgun Gothic';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Malgun Gothic')
    st.paragraph_format.line_spacing=1.1;st.paragraph_format.space_after=Pt(7)
for kind,value in blocks:
    if kind=='page':doc.add_page_break()
    elif kind=='image':
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(value[0],width=Pt(value[1]),height=Pt(value[2]))
    elif kind=='table':
        rows,widths=value;t=doc.add_table(rows=len(rows),cols=len(rows[0]));t.style='Table Grid';t.autofit=False
        for j,w in enumerate(widths):t.columns[j].width=Pt(w)
        for i,row in enumerate(rows):
            for j,v in enumerate(row):
                c=t.cell(i,j);c.width=Pt(widths[j]);c.text=str(v)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EEF3');c._tc.get_or_add_tcPr().append(shade)
                for p in c.paragraphs:
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                    for run in p.runs:run.font.size=Pt(8);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    else:
        p=doc.add_paragraph(value,{'title':'Title','h1':'Heading 1','h2':'Heading 2'}.get(kind,'Normal'))
        if kind=='small':
            for run in p.runs:run.font.size=Pt(8)
doc.core_properties.author='오성빈';doc.core_properties.title='Step2 3주차 결과보고서'
doc.save(OUT/'Step2_week3_report_OhSeongbin.docx')

pdfmetrics.registerFont(TTFont('Korean','C:/Windows/Fonts/malgun.ttf'))
pdfmetrics.registerFont(TTFont('KoreanBold','C:/Windows/Fonts/malgunbd.ttf'))
styles={k:ParagraphStyle(k,fontName='KoreanBold' if k in ('title','h1','h2') else 'Korean',
                         fontSize=size,leading=lead,spaceAfter=gap,wordWrap='CJK',textColor=colors.black)
        for k,size,lead,gap in [('title',24,33,13),('h1',16,23,12),('h2',11,18,7),('body',9.5,15,8),('small',8,12,6)]}
story=[]
for kind,value in blocks:
    if kind=='page':story.append(PageBreak())
    elif kind=='image':story.append(Image(value[0],width=value[1],height=value[2]))
    elif kind=='table':
        rows,widths=value
        t=Table([[Paragraph(escape(str(v)),styles['small']) for v in row] for row in rows],colWidths=widths,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e4eef3')),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#d9d9d9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        story.extend([t,Spacer(1,8)])
    else:story.append(Paragraph(escape(value),styles[kind]))
def footer(c,doc):
    c.setFont('Korean',8);c.drawString(44,25,'Step2 3주차 | 오성빈');c.drawRightString(551,25,str(doc.page))
pdf=OUT/'Step2_week3_report_OhSeongbin.pdf'
SimpleDocTemplate(str(pdf),pagesize=(595.28,841.89),leftMargin=44,rightMargin=44,topMargin=40,bottomMargin=48,title='Step2 3주차 결과보고서',author='오성빈').build(story,onFirstPage=footer,onLaterPages=footer)
qa=ROOT/'tmp/week3_report_qa';qa.mkdir(parents=True,exist_ok=True)
rendered=pypdfium2.PdfDocument(str(pdf))
for i in range(len(rendered)):rendered[i].render(scale=1.4).to_pil().save(qa/f'page_{i+1}.png')
print('Created PDF pages:',len(rendered))
print(pdf)
